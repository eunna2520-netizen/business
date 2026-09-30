"""data/weather.json + data/template.xlsx -> 공기산정_2026기준.xlsx"""
import json

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

import sys
conds = json.load(open(sys.argv[1] if len(sys.argv) > 1 else "data/weather.json", encoding="utf-8"))
import sys
wb = openpyxl.load_workbook("data/template.xlsx")

# 1) 제주도 지점 데이터 시트
ds = wb.create_sheet("기상자료")
ds.append(["지점코드", "행정구역", "지점명", "조건번호", "조건", *[f"{m}월" for m in range(1, 13)], "소계"])
KEEP = ["184", "189", "185", "188"]  # 제주, 서귀포, 고산, 성산 (제주도 지점만 사용)
codes = [c for c in KEEP if c in conds[0]["stations"]]
assert len(codes) == 4, codes
for code in codes:
    for i, c in enumerate(conds, 1):
        s = c["stations"][code]
        ds.append([int(code), s["region"], s["name"], i, c["name"], *s["months"], s["total"]])
for cell in ds[1]:
    cell.font = Font(bold=True)
ds.freeze_panes = "A2"
ds.column_dimensions["E"].width = 34
n = ds.max_row

ls = wb.create_sheet("지점목록")
ls.append(["지점명", "행정구역", "지점코드"])
for code in codes:
    s = conds[0]["stations"][code]
    ls.append([s["name"], s["region"], int(code)])
for cell in ls[1]:
    cell.font = Font(bold=True)
np_ = ls.max_row

# 2) 별표2 시트를 수식으로 전환
wx = wb["별표2-비작업일수"]
wx["A1"] = "별표 2. 기상상태에 따른 지역별 비작업일수(2015년-2024년)  ※ 출처: 2026년 적정 공사기간 확보를 위한 가이드라인 부록 3"
wx["A2"], wx["D2"] = "혹서기 기준", "체감온도"
wx["A2"].font = wx["D2"].font = Font(bold=True)
dv = DataValidation(type="list", formula1='"체감온도,일최고기온"', allow_blank=False)
wx.add_data_validation(dv); dv.add("D2")
wx["E2"] = "← 혹서기를 체감온도/일최고기온 중 무엇으로 볼지 선택 (가이드라인 부록3 (1)(2) 또는 (3)(4))"

hot = lambda a, b: f'IF($D$2="체감온도",{a},{b})'
slot_cond = [hot(1, 3), hot(2, 4), 5, 6, 7, 8, 10, 11, 12, 13]
slot_label = [
    f'=IF($D$2="체감온도","혹서기\n체감온도 33°C 이상","혹서기\n일최고기온 33°C 이상")',
    f'=IF($D$2="체감온도","혹서기\n체감온도 35°C 이상","혹서기\n일최고기온 35°C 이상")',
    "동절기\n일최고기온 0°C 이하", "동절기\n신적설 5㎝ 이상",
    "동절기 : 일최고기온 0°C 이하 및 신적설 5㎝ 이상", "동절기 : 일최저기온 -12°C 이하",
    "일강수량 : 5㎜ 이상", "일강수량 : 10㎜ 이상", "일강수량 : 20㎜ 이상",
    "일최대순간풍속 : 15m/s 이상",
]
from copy import copy
def copy_block(ws, src_top, src_bottom, shift, c1, c2, header_rows=2):
    """서식과 병합, 행 높이를 shift행 아래로 복사 (머리글 값만 복사)."""
    for r in range(src_top, src_bottom + 1):
        ws.row_dimensions[r + shift].height = ws.row_dimensions[r].height
        for c in range(c1, c2 + 1):
            src = ws.cell(r, c)
            dst = ws.cell(r + shift, c)
            dst._style = copy(src._style)
            if r < src_top + header_rows:
                dst.value = src.value
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= src_top and m.max_row <= src_bottom and m.min_col >= c1 and m.max_col <= c2:
            ws.merge_cells(start_row=m.min_row + shift, end_row=m.max_row + shift,
                           start_column=m.min_col, end_column=m.max_col)

copy_block(wx, 17, 29, 14, 1, 18)
copy_block(wx, 17, 29, 28, 1, 18)
for top in (31, 45):  # 새 블록의 '적용' 행 (원본 29행 형식)
    a, b = top + 2, top + 11
    wx.cell(b + 1, 3).value = "적  용"
    for col in range(4, 16):
        cl = openpyxl.utils.get_column_letter(col)
        wx.cell(b + 1, col).value = f"=SUM({cl}{a}:{cl}{b})"
    wx.cell(b + 1, 16).value = f"=SUM(D{b + 1}:O{b + 1})"
wx["R3"] = "조건번호(자동)"
BLOCKS = ((5, "제주"), (19, "서귀포"), (33, "고산"), (47, "성산"))
for first, default in BLOCKS:
    last = first + 9
    code_cell = f"$A${first}"
    wx.cell(first, 2).value = default
    wx.cell(first, 1).value = f"=INDEX(지점목록!$C$2:$C${np_},MATCH($B${first},지점목록!$A$2:$A${np_},0))"
    d = DataValidation(type="list", formula1=f"=지점목록!$A$2:$A${np_}", allow_blank=False)
    wx.add_data_validation(d); d.add(f"B{first}")
    for k in range(10):
        r = first + k
        wx.cell(r, 3).value = "2015~2024"
        wx.cell(r, 18).value = f"={slot_cond[k]}" if isinstance(slot_cond[k], str) else slot_cond[k]
        wx.cell(r, 17).value = slot_label[k]
        wx.cell(r, 17).alignment = Alignment(wrap_text=True, vertical="center")
        for m in range(12):
            col = openpyxl.utils.get_column_letter(6 + m)  # 기상자료 시트의 1~12월 열(F~Q)
            wx.cell(r, 4 + m).value = (
                f"=SUMIFS(기상자료!${col}$2:${col}${n},기상자료!$A$2:$A${n},{code_cell},기상자료!$D$2:$D${n},$R{r})")
        wx.cell(r, 16).value = f"=SUM(D{r}:O{r})"
for hr in (3, 17, 31, 45):
    wx.cell(hr, 1).value = "지점코드"
    wx.cell(hr, 2).value = "지역 (선택)"
wx.column_dimensions["R"].width = 14

# 3) 공기산정 시트 연결
main = wb["공기산정"]
main["AN21"] = "=SUM(AB21:AM21)"
X1, X2 = 24, 40  # X~AN
copy_block(main, 25, 31, 16, X1, X2)   # 표3: 41~47행
copy_block(main, 25, 31, 24, X1, X2)   # 표4: 49~55행
WX = "'별표2-비작업일수'!"
for top, r0 in ((17, 5), (25, 19), (41, 33), (49, 47)):
    d0 = top + 2                        # 데이터 첫 행
    main.cell(top, 27).value = "구분"    # AA
    main.cell(d0, 25).value = f"={WX}A{r0}"    # 번호 -> 지점코드
    main.cell(d0, 26).value = f"={WX}B{r0}"    # 지역
    for k, (lab, off) in enumerate((("혹서기", 1), ("동절기", 3), ("강   우", 8), ("바람", 9))):
        rr = d0 + k
        main.cell(rr, 27).value = lab
        for m in range(12):
            main.cell(rr, 28 + m).value = f"={WX}{openpyxl.utils.get_column_letter(4 + m)}{r0 + off}"
        main.cell(rr, 40).value = f"=SUM(AB{rr}:AM{rr})"
    main.cell(d0 + 4, 27).value = "적  용"
    for m in range(12):
        cl = openpyxl.utils.get_column_letter(28 + m)
        main.cell(d0 + 4, 28 + m).value = f"=SUM({cl}{d0}:{cl}{d0 + 3})"
    main.cell(d0 + 4, 40).value = f"=SUM(AB{d0 + 4}:AM{d0 + 4})"
# 준비기간 예시를 2026 가이드라인 표로 갱신 (6행으로 늘림)
import sheets_civil as _sc
copy_block(main, 13, 13, 1, 4, 12, header_rows=0)
for i, (a1, a2, b1, b2) in enumerate(_sc.PREP_EXAMPLES):
    r = 9 + i
    main.cell(r, 4).value, main.cell(r, 7).value = a1, a2
    main.cell(r, 9).value, main.cell(r, 11).value = b1, b2
# 왼쪽 월별표(F열)는 첫 번째 기상표(19~22행)의 해당 월 값을 참조 (원본은 다른 행을 가리켜 주석과 불일치)
for r in range(19, 23):
    main.cell(r, 6).value = f'=INDEX($AB{r}:$AM{r},VALUE(SUBSTITUTE(F$17,"월","")))'

import sheets_civil
n_hol = sheets_civil.add_holiday_sheet(wb)
n_rate, n_cat = sheets_civil.add_workrate_sheet(wb)
n_types = sheets_civil.add_preset_sheet(wb)
sheets_civil.add_facility_sheet(wb)
_ws, info = sheets_civil.add_workday_sheet(wb, n, np_, n_hol, n_rate, n_types, n_cat)
sheets_civil.add_schedule_sheet(wb, info, n_hol)
sheets_civil.add_limit_reference_sheet(wb)
wb.calculation.fullCalcOnLoad = True
wb.save("공기산정_2026기준.xlsx")
print("saved", n - 1, "data rows")
