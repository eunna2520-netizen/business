import io
import os
from datetime import date, datetime

import openpyxl
from flask import Flask, render_template, request, send_file

import weather

app = Flask(__name__)
TEMPLATE = os.path.join(os.path.dirname(__file__), "data", "template.xlsx")
SHEET_MAIN, SHEET_WX = "공기산정", "별표2-비작업일수"
# 별표2 시트의 지역 블록: (첫 행, ASOS 지점번호)
BLOCKS = {"서귀포": (5, "189"), "제주": (19, "184")}


def build_workbook(tables, end_year, start, end):
    wb = openpyxl.load_workbook(TEMPLATE)
    wx = wb[SHEET_WX]
    for name, table in tables.items():
        first = BLOCKS[name][0]
        for i, row in enumerate(table):
            r = first + i
            wx.cell(r, 3).value = end_year
            for m, v in enumerate(row):
                wx.cell(r, 4 + m).value = v
            wx.cell(r, 16).value = f"=SUM(D{r}:O{r})"
    main = wb[SHEET_MAIN]
    main["AN21"] = "=SUM(AB21:AM21)"
    main["G3"], main["J3"] = start, end
    wb.calculation.fullCalcOnLoad = True
    return wb


@app.get("/")
def index():
    y = date.today().year
    return render_template("index.html", default_end=y - 1, default_years=10)


@app.post("/generate")
def generate():
    api_key = os.environ.get("WEATHER_API_KEY", "").strip()
    if not api_key:
        return "서버에 WEATHER_API_KEY가 설정되지 않았습니다.", 500
    years = max(1, min(int(request.form["years"]), 30))
    end_year = int(request.form["end_year"])
    start = datetime.strptime(request.form["start"], "%Y-%m-%d")
    end = datetime.strptime(request.form["end"], "%Y-%m-%d")
    tables = {}
    for name, (_, stn) in BLOCKS.items():
        items = []
        for y in range(end_year - years + 1, end_year + 1):
            items += weather.fetch_daily(api_key, stn, y)
        tables[name] = weather.monthly_averages(items, years)
    buf = io.BytesIO()
    build_workbook(tables, end_year, start, end).save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="공기산정.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
